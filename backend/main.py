from contextlib import asynccontextmanager
from datetime import datetime
from pathlib import Path
import re
from typing import List, Optional
from uuid import uuid4

from openpyxl import load_workbook
from fastapi import Depends, FastAPI, HTTPException, Query, status
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, ConfigDict, Field, field_validator
from sqlalchemy import (
    Boolean,
    Column,
    Float,
    ForeignKey,
    Integer,
    String,
    Text,
    create_engine,
    inspect,
    text,
)
from sqlalchemy.orm import Session, declarative_base, sessionmaker


# Keep the database path stable when starting Uvicorn from the backend folder.
SQLALCHEMY_DATABASE_URL = "sqlite:///./nexusmarket.db"
engine = create_engine(
    SQLALCHEMY_DATABASE_URL, connect_args={"check_same_thread": False}
)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()


class Business(Base):
    __tablename__ = "businesses"

    id = Column(Integer, primary_key=True, index=True, autoincrement=True)
    directory_id = Column(String, nullable=True)
    org_name = Column(String, unique=True, index=True, nullable=False)
    tag = Column(String, nullable=False)
    category = Column(String, nullable=False)
    verified = Column(Boolean, nullable=False, default=False)
    handles = Column(String, nullable=False, default="")
    meeting_time = Column(String, nullable=True)
    meeting_location = Column(String, nullable=True)


class ProductModel(Base):
    __tablename__ = "products"

    id = Column(Integer, primary_key=True, index=True, autoincrement=True)
    directory_offer_id = Column(String, nullable=True)
    business_id = Column(
        Integer, ForeignKey("businesses.id"), nullable=False
    )
    title = Column(String, index=True, nullable=False)
    offering_details = Column(Text, nullable=True)
    price_display = Column(String, nullable=True)
    promo_details = Column(Text, nullable=True)
    category = Column(String, index=True, nullable=False)
    price = Column(Float, nullable=False)
    rating = Column(Float, default=5.0)
    reviewsCount = Column(Integer, default=1)
    description = Column(Text, nullable=False)
    image = Column(String, nullable=False)
    seller = Column(String, nullable=False)
    stock = Column(Integer, default=10)
    shippingInfo = Column(String, default="Free Shipping")


class OrderModel(Base):
    __tablename__ = "orders"

    id = Column(String, primary_key=True, index=True)
    product_id = Column(Integer, ForeignKey("products.id"), nullable=True)
    productTitle = Column(String, nullable=False)
    quantity = Column(Integer, nullable=False)
    total = Column(Float, nullable=False)
    promo_code = Column(String, nullable=True)
    discount = Column(Float, nullable=False, default=0.0)
    status = Column(String, default="Processing")
    buyerName = Column(String, nullable=False)
    pau_email = Column(String, nullable=True)
    date = Column(String, nullable=False)
    shippingAddress = Column(String, nullable=False)


class PromoCode(Base):
    __tablename__ = "promo_codes"

    id = Column(Integer, primary_key=True, index=True, autoincrement=True)
    code = Column(String, unique=True, index=True, nullable=False)
    business_id = Column(Integer, ForeignKey("businesses.id"), nullable=False)
    discount_percent = Column(Float, nullable=False)
    is_active = Column(Boolean, nullable=False, default=True)


# create_all creates new tables; the migration below adds fields to existing tables.
Base.metadata.create_all(bind=engine)

WORKBOOK_PATH = Path(__file__).with_name("Campus_Platform_Discovery_Data_v2.xlsx")
CLUB_MEETING_DETAILS = {
    "meeting_time": "7pm",
    "meeting_location": "Abuja Hall",
}
OFFERING_IMAGES = {
    "Fashion, Beauty & Lifestyle": "https://images.unsplash.com/photo-1483985988355-763728e1935b?w=800&auto=format&fit=crop&q=80",
    "Tech, Gadgets & Digital Services": "https://images.unsplash.com/photo-1519389950473-47ba0277781c?w=800&auto=format&fit=crop&q=80",
    "Daily Essentials & Convenience": "https://images.unsplash.com/photo-1542838132-92c53300491e?w=800&auto=format&fit=crop&q=80",
    "Academics & Student Life": "https://images.unsplash.com/photo-1509062522246-3755977927d7?w=800&auto=format&fit=crop&q=80",
}


def read_workbook_data():
    """Read the directory and offering worksheets using their header names."""
    if not WORKBOOK_PATH.is_file():
        raise FileNotFoundError(
            f"Required campus directory workbook was not found: {WORKBOOK_PATH}"
        )

    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        required_sheets = {"Platform Directory", "Products & Offerings"}
        missing_sheets = required_sheets.difference(workbook.sheetnames)
        if missing_sheets:
            raise ValueError(
                "Workbook is missing required worksheet(s): "
                + ", ".join(sorted(missing_sheets))
            )

        def records(sheet_name):
            rows = workbook[sheet_name].iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else "" for value in next(rows)]
            for row in rows:
                values = dict(zip(headers, row))
                if any(value is not None and str(value).strip() for value in row):
                    yield values

        directory = list(records("Platform Directory"))
        offerings = list(records("Products & Offerings"))
        return directory, offerings
    finally:
        workbook.close()


def parse_starting_price(price_text: str) -> float:
    """Extract the first amount from a workbook price or price range."""
    amounts = re.findall(r"\d[\d,]*(?:\.\d+)?", price_text)
    if not amounts:
        raise ValueError(f"Could not parse offering price: {price_text!r}")
    return float(amounts[0].replace(",", ""))


def parse_verified(value) -> bool:
    return str(value or "").strip().lower() in {"yes", "true", "1", "verified"}


def migrate_legacy_database() -> None:
    """Add workbook import fields to existing SQLite databases."""
    inspector = inspect(engine)
    business_columns = {column["name"] for column in inspector.get_columns("businesses")}
    product_columns = {
        column["name"] for column in inspector.get_columns("products")
    }
    order_columns = {column["name"] for column in inspector.get_columns("orders")}
    with engine.begin() as connection:
        for name in ("directory_id", "meeting_time", "meeting_location"):
            if name not in business_columns:
                connection.execute(
                    text(f"ALTER TABLE businesses ADD COLUMN {name} VARCHAR")
                )
        connection.execute(
            text(
                "CREATE UNIQUE INDEX IF NOT EXISTS ix_businesses_directory_id "
                "ON businesses (directory_id)"
            )
        )

        if "business_id" not in product_columns:
            connection.execute(
                text(
                    "ALTER TABLE products ADD COLUMN business_id INTEGER "
                    "REFERENCES businesses(id)"
                )
            )
        connection.execute(
            text(
                "CREATE INDEX IF NOT EXISTS ix_products_business_id "
                "ON products (business_id)"
            )
        )
        for name, column_type in (
            ("directory_offer_id", "VARCHAR"),
            ("offering_details", "TEXT"),
            ("price_display", "VARCHAR"),
            ("promo_details", "TEXT"),
        ):
            if name not in product_columns:
                connection.execute(
                    text(f"ALTER TABLE products ADD COLUMN {name} {column_type}")
                )
        connection.execute(
            text(
                "CREATE UNIQUE INDEX IF NOT EXISTS "
                "uq_products_directory_offer_id ON products (directory_offer_id)"
            )
        )

        if "product_id" not in order_columns:
            connection.execute(
                text(
                    "ALTER TABLE orders ADD COLUMN product_id INTEGER "
                    "REFERENCES products(id)"
                )
            )
        if "pau_email" not in order_columns:
            connection.execute(text("ALTER TABLE orders ADD COLUMN pau_email VARCHAR"))
        if "promo_code" not in order_columns:
            connection.execute(text("ALTER TABLE orders ADD COLUMN promo_code VARCHAR"))
        if "discount" not in order_columns:
            connection.execute(
                text("ALTER TABLE orders ADD COLUMN discount FLOAT NOT NULL DEFAULT 0")
            )


def seed_database() -> None:
    """Idempotently import directory and offering rows from the supplied workbook."""
    directory_rows, offering_rows = read_workbook_data()
    db = SessionLocal()
    try:
        businesses_by_id = {
            business.directory_id: business
            for business in db.query(Business).all()
            if business.directory_id
        }
        businesses_by_name = {
            business.org_name.strip().casefold(): business
            for business in db.query(Business).all()
        }
        legacy_names = {
            "tekk business": "Tekk",
            "preetythings.co": "Pretty Things.Co",
            "dapper steeze biz": "Dapper Steeze",
        }

        for row in directory_rows:
            directory_id = str(row.get("ID") or "").strip()
            org_name = str(row.get("Organization Name") or "").strip()
            tag = str(row.get("Tag") or "").strip()
            category = str(row.get("Category") or "").strip()
            handles = str(row.get("Numbers/Handles") or "").strip()
            if not directory_id or not org_name or not tag or not category:
                raise ValueError(f"Invalid Platform Directory record: {row!r}")

            business = businesses_by_id.get(directory_id)
            if business is None:
                lookup_name = legacy_names.get(org_name.casefold(), org_name)
                business = businesses_by_name.get(lookup_name.casefold())
            if business is None:
                business = Business()
                db.add(business)

            is_club = tag.casefold() == "club"
            business.directory_id = directory_id
            business.org_name = org_name
            business.tag = tag
            business.category = category
            business.verified = parse_verified(row.get("Verified"))
            business.handles = handles
            business.meeting_time = (
                CLUB_MEETING_DETAILS["meeting_time"] if is_club else None
            )
            business.meeting_location = (
                CLUB_MEETING_DETAILS["meeting_location"] if is_club else None
            )
            businesses_by_id[directory_id] = business
            businesses_by_name[org_name.casefold()] = business
        db.commit()

        for row in offering_rows:
            directory_id = str(row.get("ID") or "").strip()
            org_name = str(row.get("Organization Name") or "").strip()
            items = str(row.get("Item Being Sold / Offered") or "").strip()
            price_display = str(row.get("Price") or "").strip()
            promo_details = str(row.get("Promos & Discounts") or "").strip()
            if not directory_id or not items or not price_display:
                raise ValueError(f"Invalid Products & Offerings record: {row!r}")

            business = businesses_by_id.get(directory_id)
            if business is None:
                business = businesses_by_name.get(org_name.casefold())
            if business is None:
                raise ValueError(
                    f"Offering {directory_id} ({org_name}) has no directory business"
                )

            product = (
                db.query(ProductModel)
                .filter(ProductModel.directory_offer_id == directory_id)
                .first()
            )
            if product is None:
                product = ProductModel()
                db.add(product)

            product.directory_offer_id = directory_id
            product.business_id = business.id
            product.title = f"{business.org_name} offerings"
            product.offering_details = items
            product.category = business.category
            product.price = parse_starting_price(price_display)
            product.price_display = price_display
            product.promo_details = promo_details or None
            product.rating = 5.0
            product.reviewsCount = 1
            product.description = items
            product.image = OFFERING_IMAGES.get(
                business.category,
                "https://images.unsplash.com/photo-1523275335684-37898b6baf30?w=800&auto=format&fit=crop&q=80",
            )
            product.seller = business.org_name
            product.stock = 10
            product.shippingInfo = "Contact vendor to arrange your order"
        db.commit()

        # These starter codes are scoped to individual vendors and can be changed
        # or disabled by updating the corresponding promo_codes rows.
        starter_promos = {
            "TEKK10": "Tekk",
            "PREETY10": "Pretty Things.Co",
            "DAPPER10": "Dapper Steeze",
            "JOZIE10": "Jozie Glam and Glow",
        }
        for code, org_name in starter_promos.items():
            business = (
                db.query(Business)
                .filter(Business.org_name.ilike(org_name))
                .first()
            )
            if business is None:
                continue
            promo = (
                db.query(PromoCode)
                .filter(PromoCode.code == code)
                .first()
            )
            if promo is None:
                db.add(
                    PromoCode(
                        code=code,
                        business_id=business.id,
                        discount_percent=10,
                        is_active=True,
                    )
                )
        db.commit()
    finally:
        db.close()


@asynccontextmanager
async def lifespan(app: FastAPI):
    migrate_legacy_database()
    seed_database()
    yield


app = FastAPI(
    title="PAU Connect API",
    description="Campus business discovery and booking API.",
    version="2.0.0",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


class BusinessResponse(BaseModel):
    id: int
    directory_id: Optional[str]
    org_name: str
    tag: str
    category: str
    verified: bool
    handles: str
    meeting_time: Optional[str]
    meeting_location: Optional[str]

    model_config = ConfigDict(from_attributes=True)


class ProductCreate(BaseModel):
    business_id: int
    title: str
    category: str
    price: float = Field(gt=0)
    description: str
    offering_details: Optional[str] = None
    price_display: Optional[str] = None
    promo_details: Optional[str] = None
    image: Optional[str] = (
        "https://images.unsplash.com/photo-1523275335684-37898b6baf30?w=600&auto=format&fit=crop&q=80"
    )
    seller: Optional[str] = None
    stock: int = Field(default=10, ge=0)
    shippingInfo: Optional[str] = "Campus collection"


class ProductResponse(ProductCreate):
    id: int
    directory_offer_id: Optional[str]
    rating: float
    reviewsCount: int

    model_config = ConfigDict(from_attributes=True)


class OrderLine(BaseModel):
    product_id: int
    quantity: int = Field(gt=0)


class OrderCreate(BaseModel):
    pau_email: str
    items: List[OrderLine] = Field(min_length=1)
    promo_code: Optional[str] = Field(default=None, max_length=32)

    @field_validator("pau_email")
    @classmethod
    def validate_pau_email(cls, value: str) -> str:
        if not value.endswith("@pau.edu.ng") or value.count("@") != 1:
            raise ValueError("Email must end with @pau.edu.ng")
        local_part = value.split("@", 1)[0]
        if not local_part or any(character.isspace() for character in value):
            raise ValueError("Enter a valid @pau.edu.ng email address")
        return value

    @field_validator("promo_code")
    @classmethod
    def normalize_promo_code(cls, value: Optional[str]) -> Optional[str]:
        if value is None:
            return None
        normalized = value.strip().upper()
        if not normalized:
            return None
        if not re.fullmatch(r"[A-Z0-9-]+", normalized):
            raise ValueError("Enter a valid promo code")
        return normalized


class PromoCodeCheck(BaseModel):
    code: str = Field(min_length=1, max_length=32)
    items: List[OrderLine] = Field(min_length=1)

    @field_validator("code")
    @classmethod
    def normalize_code(cls, value: str) -> str:
        normalized = value.strip().upper()
        if not re.fullmatch(r"[A-Z0-9-]+", normalized):
            raise ValueError("Enter a valid promo code")
        return normalized


class PromoCodeResponse(BaseModel):
    code: str
    business_id: int
    business_name: str
    discount_percent: float


class PromoCodeCheckResponse(PromoCodeResponse):
    eligible_subtotal: float
    discount_amount: float


class OrderResponse(BaseModel):
    id: str
    product_id: Optional[int]
    productTitle: str
    quantity: int
    total: float
    promo_code: Optional[str]
    discount: float
    status: str
    buyerName: str
    pau_email: Optional[str]
    date: str
    shippingAddress: str

    model_config = ConfigDict(from_attributes=True)


class OrderStatusUpdate(BaseModel):
    status: str


@app.get("/")
def read_root():
    return {"message": "Welcome to PAU Connect API", "docs": "/docs"}


@app.get("/api/businesses", response_model=List[BusinessResponse])
def get_businesses(db: Session = Depends(get_db)):
    return (
        db.query(Business)
        .filter(Business.directory_id.is_not(None))
        .order_by(Business.org_name)
        .all()
    )


@app.get("/api/businesses/{business_id}", response_model=BusinessResponse)
def get_business(business_id: int, db: Session = Depends(get_db)):
    business = db.query(Business).filter(Business.id == business_id).first()
    if business is None:
        raise HTTPException(status_code=404, detail="Business not found")
    return business


@app.get("/api/businesses/{business_id}/products", response_model=List[ProductResponse])
def get_business_products(business_id: int, db: Session = Depends(get_db)):
    business_exists = (
        db.query(Business.id).filter(Business.id == business_id).first()
    )
    if business_exists is None:
        raise HTTPException(status_code=404, detail="Business not found")
    return (
        db.query(ProductModel)
        .filter(
            ProductModel.business_id == business_id,
            ProductModel.directory_offer_id.is_not(None),
        )
        .order_by(ProductModel.id)
        .all()
    )


@app.get("/api/products", response_model=List[ProductResponse])
def get_products(
    category: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    business_id: Optional[int] = Query(None, gt=0),
    db: Session = Depends(get_db),
):
    # Keep legacy prototype products out of the workbook-backed marketplace.
    query = db.query(ProductModel).filter(
        ProductModel.directory_offer_id.is_not(None)
    )
    if business_id is not None:
        query = query.filter(
            ProductModel.business_id == business_id,
            ProductModel.directory_offer_id.is_not(None),
        )
    if category and category != "All Categories":
        query = query.filter(ProductModel.category == category)
    if search:
        search_filter = f"%{search}%"
        query = query.filter(
            (ProductModel.title.ilike(search_filter))
            | (ProductModel.description.ilike(search_filter))
            | (ProductModel.seller.ilike(search_filter))
        )
    return query.order_by(ProductModel.id).all()


def get_promo_discount(
    db: Session,
    code: str,
    products_by_id: dict,
    quantities_by_id: dict,
):
    promo = (
        db.query(PromoCode)
        .filter(PromoCode.code == code, PromoCode.is_active.is_(True))
        .first()
    )
    if promo is None:
        raise HTTPException(status_code=404, detail="That promo code is not valid.")

    business = db.query(Business).filter(Business.id == promo.business_id).first()
    if business is None:
        raise HTTPException(
            status_code=409,
            detail="This promo code is not linked to an available business.",
        )
    if not 0 < promo.discount_percent <= 100:
        raise HTTPException(
            status_code=409,
            detail="This promo code has an invalid discount configuration.",
        )
    eligible_subtotal = sum(
        products_by_id[product_id].price * quantity
        for product_id, quantity in quantities_by_id.items()
        if products_by_id[product_id].business_id == promo.business_id
    )
    if eligible_subtotal <= 0:
        raise HTTPException(
            status_code=422,
            detail=f"{code} is only valid for products from {business.org_name}.",
        )

    return promo, business, eligible_subtotal, round(
        eligible_subtotal * promo.discount_percent / 100, 2
    )


@app.get("/api/promo-codes", response_model=List[PromoCodeResponse])
def get_promo_codes(
    business_id: Optional[int] = Query(None, gt=0),
    db: Session = Depends(get_db),
):
    query = db.query(PromoCode, Business).join(
        Business, PromoCode.business_id == Business.id
    )
    query = query.filter(PromoCode.is_active.is_(True))
    if business_id is not None:
        query = query.filter(PromoCode.business_id == business_id)
    return [
        {
            "code": promo.code,
            "business_id": promo.business_id,
            "business_name": business.org_name,
            "discount_percent": promo.discount_percent,
        }
        for promo, business in query.order_by(PromoCode.code).all()
    ]


@app.post("/api/promo-codes/validate", response_model=PromoCodeCheckResponse)
def validate_promo_code(payload: PromoCodeCheck, db: Session = Depends(get_db)):
    quantities_by_id = {}
    for item in payload.items:
        quantities_by_id[item.product_id] = (
            quantities_by_id.get(item.product_id, 0) + item.quantity
        )
    products_by_id = {}
    for product_id in quantities_by_id:
        product = db.query(ProductModel).filter(ProductModel.id == product_id).first()
        if product is None:
            raise HTTPException(status_code=404, detail=f"Product {product_id} not found")
        products_by_id[product_id] = product

    promo, business, subtotal, discount = get_promo_discount(
        db, payload.code, products_by_id, quantities_by_id
    )
    return {
        "code": promo.code,
        "business_id": business.id,
        "business_name": business.org_name,
        "discount_percent": promo.discount_percent,
        "eligible_subtotal": round(subtotal, 2),
        "discount_amount": discount,
    }


@app.post(
    "/api/products",
    response_model=ProductResponse,
    status_code=status.HTTP_201_CREATED,
)
def create_product(product: ProductCreate, db: Session = Depends(get_db)):
    business = (
        db.query(Business)
        .filter(Business.id == product.business_id)
        .first()
    )
    if business is None:
        raise HTTPException(status_code=404, detail="Business not found")
    db_product = ProductModel(
        **product.model_dump(exclude={"seller"}),
        seller=product.seller or business.org_name,
        rating=5.0,
        reviewsCount=1,
    )
    db.add(db_product)
    db.commit()
    db.refresh(db_product)
    return db_product


@app.delete("/api/products/{product_id}", status_code=status.HTTP_200_OK)
def delete_product(product_id: int, db: Session = Depends(get_db)):
    product = db.query(ProductModel).filter(ProductModel.id == product_id).first()
    if product is None:
        raise HTTPException(status_code=404, detail="Product not found")
    db.delete(product)
    db.commit()
    return {"message": "Product deleted successfully", "id": product_id}


@app.get("/api/orders", response_model=List[OrderResponse])
def get_orders(db: Session = Depends(get_db)):
    return db.query(OrderModel).order_by(OrderModel.date.desc()).all()


@app.post(
    "/api/orders",
    response_model=List[OrderResponse],
    status_code=status.HTTP_201_CREATED,
)
def create_order(payload: OrderCreate, db: Session = Depends(get_db)):
    quantities_by_id = {}
    for item in payload.items:
        quantities_by_id[item.product_id] = (
            quantities_by_id.get(item.product_id, 0) + item.quantity
        )

    products_by_id = {}
    for product_id, quantity in quantities_by_id.items():
        product = db.query(ProductModel).filter(ProductModel.id == product_id).first()
        if product is None:
            raise HTTPException(status_code=404, detail=f"Product {product_id} not found")
        if product.stock < quantity:
            raise HTTPException(
                status_code=409,
                detail=f"Not enough stock for {product.title}",
            )
        products_by_id[product_id] = product

    promo_details = None
    if payload.promo_code:
        promo_details = get_promo_discount(
            db,
            payload.promo_code,
            products_by_id,
            quantities_by_id,
        )

    current_date = datetime.now().strftime("%Y-%m-%d")
    created_orders = []
    eligible_product_ids = (
        [
            product_id
            for product_id, product in products_by_id.items()
            if promo_details and product.business_id == promo_details[0].business_id
        ]
        if promo_details
        else []
    )
    remaining_discount = promo_details[3] if promo_details else 0.0
    for product_id, quantity in quantities_by_id.items():
        product = products_by_id[product_id]
        product.stock -= quantity
        line_total = product.price * quantity
        line_discount = 0.0
        if product_id in eligible_product_ids:
            if len(eligible_product_ids) == 1:
                line_discount = remaining_discount
            else:
                line_discount = min(
                    remaining_discount,
                    round(
                        line_total * promo_details[0].discount_percent / 100, 2
                    ),
                )
                remaining_discount = round(remaining_discount - line_discount, 2)
            eligible_product_ids.remove(product_id)
        order = OrderModel(
            id=f"ORD-{uuid4().hex[:12].upper()}",
            product_id=product.id,
            productTitle=product.title,
            quantity=quantity,
            total=round(line_total - line_discount, 2),
            promo_code=promo_details[0].code if line_discount else None,
            discount=line_discount,
            status="Processing",
            buyerName=payload.pau_email,
            pau_email=payload.pau_email,
            date=current_date,
            shippingAddress="PAU campus collection",
        )
        db.add(order)
        created_orders.append(order)

    db.commit()
    for order in created_orders:
        db.refresh(order)
    return created_orders


@app.patch("/api/orders/{order_id}", response_model=OrderResponse)
def update_order_status(
    order_id: str, payload: OrderStatusUpdate, db: Session = Depends(get_db)
):
    if payload.status not in {"Processing", "Shipped", "Delivered", "Cancelled"}:
        raise HTTPException(status_code=422, detail="Invalid order status")
    order = db.query(OrderModel).filter(OrderModel.id == order_id).first()
    if order is None:
        raise HTTPException(status_code=404, detail="Order not found")
    order.status = payload.status
    db.commit()
    db.refresh(order)
    return order
